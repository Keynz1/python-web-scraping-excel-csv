import csv
from pathlib import Path

from openpyxl import load_workbook

try:
    import export_data as _export
except ModuleNotFoundError:
    _export = None


ROWS = [
    {"title": "Book A", "price_gbp": "12.34", "availability": "In stock", "product_url": "a"},
    {"title": "Book B", "price_gbp": "7.00", "availability": "In stock", "product_url": "b"},
]


def impl():
    assert _export is not None, "export_data module is not implemented yet"
    return _export


def read_csv_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def read_xlsx_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    values = list(ws.iter_rows(values_only=True))
    headers = [str(v) for v in values[0]]
    return [{headers[i]: "" if value is None else str(value) for i, value in enumerate(row)} for row in values[1:]]


def test_csv_and_xlsx_contain_identical_rows(tmp_path):
    csv_path = tmp_path / "out.csv"
    xlsx_path = tmp_path / "out.xlsx"
    mod = impl()
    mod.write_csv(ROWS, csv_path)
    mod.write_xlsx(ROWS, xlsx_path)
    assert read_csv_rows(csv_path) == read_xlsx_rows(xlsx_path) == ROWS


def test_report_contains_source_and_counts(tmp_path):
    report = tmp_path / "report.txt"
    impl().write_report(ROWS, report, "https://books.toscrape.com/")
    text = report.read_text(encoding="utf-8")
    assert "Source: https://books.toscrape.com/" in text
    assert "Rows exported: 2" in text
    assert "Rows with missing price: 0" in text
